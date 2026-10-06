from io import BytesIO
from datetime import datetime, timedelta

import pandas as pd
import flask_login
from flask import request, render_template, send_file

from main import app, db
from models.SolicitacaoExame_model import SolicitacaoExame, ItemSolicitacao
from models.Exame_model import Exame


# Paleta de cores usada nos gráficos (cicla se houver mais categorias que cores)
CORES_CATEGORIAS = ['#2563eb', '#0ea5e9', '#22c55e', '#f59e0b', '#a855f7', '#ef4444', '#94a3b8']

# Ajuste aqui se o status "concluído" do seu sistema tiver outro valor
STATUS_CONCLUIDO = 'concluido'


# --------------------------------------------------------------------------
# Helpers de filtro e coleta de dados
# --------------------------------------------------------------------------

def _parse_data(valor, padrao):
    if not valor:
        return padrao
    try:
        return datetime.strptime(valor, '%Y-%m-%d').date()
    except ValueError:
        return padrao


def _ler_filtros():
    """Lê os filtros da query string, com padrão de 'últimos 6 meses'."""
    hoje = datetime.utcnow().date()
    padrao_inicio = hoje - timedelta(days=180)

    return {
        'data_inicio': _parse_data(request.args.get('data_inicio'), padrao_inicio),
        'data_fim': _parse_data(request.args.get('data_fim'), hoje),
        'tipo_relatorio': request.args.get('tipo_relatorio', 'exames'),
        'tipo_exame': request.args.get('tipo_exame', ''),
    }


def _periodo_anterior(filtros):
    """Mesma duração do período atual, imediatamente anterior a ele (para calcular variação %)."""
    duracao_dias = (filtros['data_fim'] - filtros['data_inicio']).days + 1
    fim_anterior = filtros['data_inicio'] - timedelta(days=1)
    inicio_anterior = fim_anterior - timedelta(days=duracao_dias - 1)
    return {**filtros, 'data_inicio': inicio_anterior, 'data_fim': fim_anterior}


def _coletar_dados_relatorio(filtros):
    """
    Monta um DataFrame com um registro por item de solicitação concluído dentro do
    período, cruzando ItemSolicitacao + SolicitacaoExame + Exame.
    """
    query = (
        db.session.query(
            ItemSolicitacao.id.label('item_id'),
            ItemSolicitacao.valor.label('valor_item'),
            SolicitacaoExame.paciente_id.label('paciente_id'),
            SolicitacaoExame.data_agendada.label('data_agendada'),
            SolicitacaoExame.data_solicitacao.label('data_solicitacao'),
            Exame.id.label('exame_id'),
            Exame.nome.label('exame_nome'),
            Exame.categoria.label('categoria'),
            Exame.preco.label('preco_exame'),
        )
        .join(SolicitacaoExame, ItemSolicitacao.solicitacao_id == SolicitacaoExame.id)
        .join(Exame, ItemSolicitacao.exame_id == Exame.id)
        .filter(ItemSolicitacao.status == STATUS_CONCLUIDO)
        .filter(SolicitacaoExame.data_agendada.isnot(None))
        .filter(SolicitacaoExame.data_agendada >= filtros['data_inicio'])
        .filter(SolicitacaoExame.data_agendada <= filtros['data_fim'])
    )

    if filtros.get('tipo_exame'):
        query = query.filter(Exame.id == filtros['tipo_exame'])

    linhas = query.all()

    registros = []
    for l in linhas:
        valor = float(l.valor_item) if l.valor_item is not None else float(l.preco_exame or 0)

        dias_ate_agendamento = None
        if l.data_solicitacao and l.data_agendada:
            dias_ate_agendamento = (l.data_agendada - l.data_solicitacao.date()).days

        registros.append({
            'exame_id': l.exame_id,
            'exame_nome': l.exame_nome,
            'categoria': l.categoria or 'Sem categoria',
            'valor': valor,
            'paciente_id': l.paciente_id,
            'data_agendada': l.data_agendada,
            'mes': l.data_agendada.strftime('%Y-%m'),
            'dias_ate_agendamento': dias_ate_agendamento,
        })

    colunas = ['exame_id', 'exame_nome', 'categoria', 'valor', 'paciente_id',
               'data_agendada', 'mes', 'dias_ate_agendamento']
    return pd.DataFrame(registros, columns=colunas)


# --------------------------------------------------------------------------
# Helpers de formatação / agregação
# --------------------------------------------------------------------------

def _fmt_moeda(valor):
    return f"R$ {valor:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')


def _fmt_dias(valor):
    if valor is None or pd.isna(valor):
        return 'N/D'
    return f"{valor:.1f} dias".replace('.', ',')


def _variacao(atual, anterior):
    if not anterior:
        return (0.0, 'flat') if not atual else (100.0, 'up')
    pct = round(((atual - anterior) / anterior) * 100, 1)
    direcao = 'up' if pct > 0 else 'down' if pct < 0 else 'flat'
    return pct, direcao


def _montar_kpis(df, df_anterior):
    exames_atual, exames_anterior = len(df), len(df_anterior)
    pac_atual = df['paciente_id'].nunique() if not df.empty else 0
    pac_anterior = df_anterior['paciente_id'].nunique() if not df_anterior.empty else 0
    receita_atual = df['valor'].sum() if not df.empty else 0.0
    receita_anterior = df_anterior['valor'].sum() if not df_anterior.empty else 0.0

    tempo_atual = df['dias_ate_agendamento'].mean() if not df.empty else None
    tempo_anterior = df_anterior['dias_ate_agendamento'].mean() if not df_anterior.empty else None

    var_exames, dir_exames = _variacao(exames_atual, exames_anterior)
    var_pac, dir_pac = _variacao(pac_atual, pac_anterior)
    var_receita, dir_receita = _variacao(receita_atual, receita_anterior)

    if tempo_atual is not None and tempo_anterior:
        var_tempo, dir_tempo = _variacao(tempo_atual, tempo_anterior)
    else:
        var_tempo, dir_tempo = 0.0, 'flat'

    return {
        'exames_realizados': {'valor': exames_atual, 'variacao': var_exames, 'direcao': dir_exames},
        'pacientes_atendidos': {'valor': pac_atual, 'variacao': var_pac, 'direcao': dir_pac},
        'tempo_medio_entrega': {'valor': _fmt_dias(tempo_atual), 'variacao': var_tempo, 'direcao': dir_tempo},
        'receita_total': {'valor': _fmt_moeda(receita_atual), 'variacao': var_receita, 'direcao': dir_receita},
    }


MESES_PT = {'01': 'Jan', '02': 'Fev', '03': 'Mar', '04': 'Abr', '05': 'Mai', '06': 'Jun',
            '07': 'Jul', '08': 'Ago', '09': 'Set', '10': 'Out', '11': 'Nov', '12': 'Dez'}


def _grafico_mensal(df):
    if df.empty:
        return {'labels': [], 'dados': []}
    agrupado = df.groupby('mes').size().sort_index()
    labels = [f"{MESES_PT[m.split('-')[1]]}/{m.split('-')[0][2:]}" for m in agrupado.index]
    return {'labels': labels, 'dados': agrupado.tolist()}


def _grafico_e_detalhamento_por_categoria(df):
    if df.empty:
        return {'labels': [], 'dados': [], 'cores': []}, []

    agrupado = (
        df.groupby('categoria')
        .agg(quantidade=('exame_id', 'count'), receita=('valor', 'sum'), tempo_medio=('dias_ate_agendamento', 'mean'))
        .sort_values('quantidade', ascending=False)
    )
    total = agrupado['quantidade'].sum()

    labels = agrupado.index.tolist()
    dados = agrupado['quantidade'].tolist()
    cores = [CORES_CATEGORIAS[i % len(CORES_CATEGORIAS)] for i in range(len(labels))]

    detalhamento = [
        {
            'tipo': categoria,
            'quantidade': int(row['quantidade']),
            'percentual': round((row['quantidade'] / total) * 100, 1) if total else 0,
            'tempo_medio': _fmt_dias(row['tempo_medio']),
            'receita': _fmt_moeda(row['receita']),
        }
        for categoria, row in agrupado.iterrows()
    ]

    return {'labels': labels, 'dados': dados, 'cores': cores}, detalhamento


def _carregar_relatorio_completo(filtros):
    df = _coletar_dados_relatorio(filtros)
    df_anterior = _coletar_dados_relatorio(_periodo_anterior(filtros))

    kpis = _montar_kpis(df, df_anterior)
    grafico_mensal = _grafico_mensal(df)
    grafico_tipos, detalhamento = _grafico_e_detalhamento_por_categoria(df)

    return kpis, grafico_mensal, grafico_tipos, detalhamento


# --------------------------------------------------------------------------
# Rotas
# --------------------------------------------------------------------------

@app.route('/managerment_reports')
@flask_login.login_required
def managerment_reports():
    filtros = _ler_filtros()
    kpis, grafico_mensal, grafico_tipos, detalhamento = _carregar_relatorio_completo(filtros)
    tipos_exame = Exame.query.order_by(Exame.nome).all()

    filtros_template = {
        'data_inicio': filtros['data_inicio'].strftime('%Y-%m-%d'),
        'data_fim': filtros['data_fim'].strftime('%Y-%m-%d'),
        'tipo_relatorio': filtros['tipo_relatorio'],
        'tipo_exame': filtros['tipo_exame'],
    }

    return render_template(
        'managerment_reports.html',
        user=flask_login.current_user,
        filtros=filtros_template,
        tipos_exame=tipos_exame,
        kpis=kpis,
        grafico_mensal=grafico_mensal,
        grafico_tipos=grafico_tipos,
        detalhamento=detalhamento,
    )


@app.route('/managerment_reports/exportar/excel')
@flask_login.login_required
def exportar_relatorio_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    filtros = _ler_filtros()
    _, _, _, detalhamento = _carregar_relatorio_completo(filtros)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Relatório Gerencial'

    ws.merge_cells('A1:E1')
    ws['A1'] = 'Relatório Gerencial - Laboratório de Análises Clínicas'
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 24

    ws.merge_cells('A2:E2')
    ws['A2'] = f"Período: {filtros['data_inicio'].strftime('%d/%m/%Y')} a {filtros['data_fim'].strftime('%d/%m/%Y')}"

    cabecalhos = ['Tipo de exame', 'Quantidade', '% do total', 'Tempo médio', 'Receita']
    for col, texto in enumerate(cabecalhos, start=1):
        celula = ws.cell(row=4, column=col, value=texto)
        celula.font = Font(bold=True)
        celula.fill = PatternFill(start_color='E5E7EB', end_color='E5E7EB', fill_type='solid')

    linha = 5
    for item in detalhamento:
        ws.cell(row=linha, column=1, value=item['tipo'])
        ws.cell(row=linha, column=2, value=item['quantidade'])
        ws.cell(row=linha, column=3, value=f"{item['percentual']}%")
        ws.cell(row=linha, column=4, value=item['tempo_medio'])
        ws.cell(row=linha, column=5, value=item['receita'])
        linha += 1

    if not detalhamento:
        ws.cell(row=linha, column=1, value='Nenhum dado no período selecionado.')

    for col in range(1, 6):
        ws.column_dimensions[get_column_letter(col)].width = 24

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"relatorio_gerencial_{filtros['data_inicio']}_{filtros['data_fim']}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/managerment_reports/exportar/pdf')
@flask_login.login_required
def exportar_relatorio_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    filtros = _ler_filtros()
    kpis, _, _, detalhamento = _carregar_relatorio_completo(filtros)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    estilos = getSampleStyleSheet()
    elementos = []

    elementos.append(Paragraph('Relatório Gerencial - Laboratório de Análises Clínicas', estilos['Title']))
    elementos.append(Paragraph(
        f"Período: {filtros['data_inicio'].strftime('%d/%m/%Y')} a {filtros['data_fim'].strftime('%d/%m/%Y')}",
        estilos['Normal'],
    ))
    elementos.append(Spacer(1, 0.5 * cm))

    dados_kpi = [
        ['Indicador', 'Valor'],
        ['Exames realizados', str(kpis['exames_realizados']['valor'])],
        ['Pacientes atendidos', str(kpis['pacientes_atendidos']['valor'])],
        ['Tempo médio de entrega', kpis['tempo_medio_entrega']['valor']],
        ['Receita total', kpis['receita_total']['valor']],
    ]
    tabela_kpi = Table(dados_kpi, colWidths=[8 * cm, 6 * cm])
    tabela_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elementos.append(tabela_kpi)
    elementos.append(Spacer(1, 1 * cm))

    elementos.append(Paragraph('Detalhamento por tipo de exame', estilos['Heading2']))
    cabecalhos = ['Tipo de exame', 'Qtd.', '% total', 'Tempo médio', 'Receita']
    linhas_det = [[d['tipo'], str(d['quantidade']), f"{d['percentual']}%", d['tempo_medio'], d['receita']]
                  for d in detalhamento]
    dados_det = [cabecalhos] + linhas_det

    tabela_det = Table(dados_det, colWidths=[5 * cm, 2 * cm, 2.2 * cm, 3 * cm, 3.5 * cm])
    tabela_det.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E5E7EB')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elementos.append(tabela_det)

    if not detalhamento:
        elementos.append(Spacer(1, 0.3 * cm))
        elementos.append(Paragraph('Nenhum dado disponível para o período selecionado.', estilos['Normal']))

    doc.build(elementos)
    buffer.seek(0)

    nome_arquivo = f"relatorio_gerencial_{filtros['data_inicio']}_{filtros['data_fim']}.pdf"
    return send_file(buffer, as_attachment=True, download_name=nome_arquivo, mimetype='application/pdf')